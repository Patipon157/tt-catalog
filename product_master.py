"""
Thailand Trophy - Product Master File (ต้นทุนสินค้า)
โปรแกรมจัดการข้อมูลสินค้าและต้นทุน สำหรับ Thailand Trophy

ความสามารถ:
1. เพิ่ม/แก้ไข/ลบ ข้อมูลสินค้า
2. ค้นหาสินค้า
3. กรองตามหมวดหมู่
4. คำนวณกำไร (ราคาขาย - ต้นทุน)
5. Export เป็น Excel สวยงาม (.xlsx) + CSV
6. Import จาก CSV
7. สรุปยอดต้นทุนรวม
"""

import os
import csv
import tkinter as tk
from tkinter import ttk, messagebox, filedialog, simpledialog
from datetime import datetime

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

# ===== ตั้งค่า =====
DATA_DIR = os.path.dirname(os.path.abspath(__file__))
MASTER_FILE = os.path.join(DATA_DIR, "product_master_data.csv")

# ===== หมวดหมู่สินค้า Thailand Trophy =====
CATEGORIES = [
    "ถ้วยรางวัล",
    "โล่รางวัล",
    "เหรียญรางวัล",
    "พวงกุญแจ",
    "เข็มกลัด/Pin",
    "ของที่ระลึก",
    "งานพิมพ์/สกรีน",
    "วัตถุดิบ",
    "บรรจุภัณฑ์",
    "อื่นๆ",
]

# ===== CSV Header =====
CSV_HEADERS = [
    "รหัสสินค้า",
    "ชื่อสินค้า",
    "หมวดหมู่",
    "ต้นทุน (บาท)",
    "ราคาขาย (บาท)",
    "Supplier",
    "จำนวนคงเหลือ",
    "หน่วย",
    "หมายเหตุ",
    "วันที่อัปเดต",
]


def load_data(filepath):
    """โหลดข้อมูลจาก CSV"""
    data = []
    if not os.path.exists(filepath):
        return data
    try:
        with open(filepath, "r", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            for row in reader:
                data.append(row)
    except Exception:
        pass
    return data


def save_data(filepath, data):
    """บันทึกข้อมูลลง CSV"""
    with open(filepath, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=CSV_HEADERS)
        writer.writeheader()
        writer.writerows(data)


def format_money(value):
    """แสดงตัวเลขเงินให้อ่านง่าย"""
    try:
        num = float(value)
        if num == int(num):
            return f"{int(num):,}"
        return f"{num:,.2f}"
    except (ValueError, TypeError):
        return str(value)


class ProductMaster:
    def __init__(self, root):
        self.root = root
        self.root.title("Thailand Trophy - Product Master File (ต้นทุนสินค้า)")
        self.root.geometry("1100x700")
        self.root.resizable(True, True)

        self.data = []
        self.setup_ui()
        self.load_master_file()

    def setup_ui(self):
        """สร้างหน้าจอ"""
        # ===== ส่วนบน: ปุ่มคำสั่ง =====
        cmd_frame = ttk.LabelFrame(self.root, text="จัดการสินค้า", padding=10)
        cmd_frame.pack(fill="x", padx=10, pady=5)

        ttk.Button(cmd_frame, text="เพิ่มสินค้า", command=self.add_product).pack(side="left", padx=5)
        ttk.Button(cmd_frame, text="แก้ไข", command=self.edit_product).pack(side="left", padx=5)
        ttk.Button(cmd_frame, text="ลบ", command=self.delete_product).pack(side="left", padx=5)
        ttk.Separator(cmd_frame, orient="vertical").pack(side="left", fill="y", padx=10)
        ttk.Button(cmd_frame, text="บันทึก", command=self.save_master_file).pack(side="left", padx=5)
        ttk.Button(cmd_frame, text="Export Excel", command=self.export_excel).pack(side="left", padx=5)
        ttk.Button(cmd_frame, text="Export CSV...", command=self.export_csv).pack(side="left", padx=5)
        ttk.Button(cmd_frame, text="Import CSV...", command=self.import_csv).pack(side="left", padx=5)

        # ===== ส่วนค้นหา/กรอง =====
        filter_frame = ttk.LabelFrame(self.root, text="ค้นหา / กรอง", padding=10)
        filter_frame.pack(fill="x", padx=10, pady=5)

        ttk.Label(filter_frame, text="ค้นหา:").pack(side="left", padx=(0, 5))
        self.search_var = tk.StringVar()
        self.search_var.trace("w", lambda *args: self.refresh_table())
        ttk.Entry(filter_frame, textvariable=self.search_var, width=25).pack(side="left", padx=2)

        ttk.Label(filter_frame, text="  หมวดหมู่:").pack(side="left", padx=(15, 5))
        self.cat_filter_var = tk.StringVar(value="ทั้งหมด")
        cat_combo = ttk.Combobox(
            filter_frame,
            textvariable=self.cat_filter_var,
            values=["ทั้งหมด"] + CATEGORIES,
            state="readonly",
            width=15,
        )
        cat_combo.pack(side="left", padx=2)
        cat_combo.bind("<<ComboboxSelected>>", lambda e: self.refresh_table())

        # ===== สรุปยอด =====
        self.summary_var = tk.StringVar(value="")
        ttk.Label(self.root, textvariable=self.summary_var, font=("", 10, "bold")).pack(
            fill="x", padx=10, pady=2
        )

        # ===== ตารางสินค้า =====
        table_frame = ttk.Frame(self.root)
        table_frame.pack(fill="both", expand=True, padx=10, pady=5)

        columns = (
            "code", "name", "category", "cost", "price",
            "profit", "margin", "supplier", "stock", "unit", "note", "updated"
        )
        self.tree = ttk.Treeview(table_frame, columns=columns, show="headings", height=18)

        col_config = [
            ("code", "รหัส", 80),
            ("name", "ชื่อสินค้า", 200),
            ("category", "หมวดหมู่", 100),
            ("cost", "ต้นทุน", 80),
            ("price", "ราคาขาย", 80),
            ("profit", "กำไร", 80),
            ("margin", "Margin%", 70),
            ("supplier", "Supplier", 100),
            ("stock", "คงเหลือ", 70),
            ("unit", "หน่วย", 50),
            ("note", "หมายเหตุ", 120),
            ("updated", "อัปเดต", 90),
        ]

        for col_id, heading, width in col_config:
            self.tree.heading(col_id, text=heading, command=lambda c=col_id: self.sort_column(c))
            self.tree.column(col_id, width=width, minwidth=40)

        scrollbar_y = ttk.Scrollbar(table_frame, orient="vertical", command=self.tree.yview)
        scrollbar_x = ttk.Scrollbar(table_frame, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=scrollbar_y.set, xscrollcommand=scrollbar_x.set)

        self.tree.pack(side="left", fill="both", expand=True)
        scrollbar_y.pack(side="right", fill="y")
        scrollbar_x.pack(side="bottom", fill="x")

        self.tree.bind("<Double-1>", lambda e: self.edit_product())

        # ===== สถานะ =====
        self.status_var = tk.StringVar(value="พร้อมใช้งาน")
        ttk.Label(self.root, textvariable=self.status_var, relief="sunken", padding=5).pack(
            fill="x", padx=10, pady=5
        )

        self.sort_reverse = {}

    # ===== Data I/O =====

    def load_master_file(self):
        """โหลดข้อมูล Master File"""
        self.data = load_data(MASTER_FILE)
        self.refresh_table()
        count = len(self.data)
        self.status_var.set(f"โหลดข้อมูล {count} รายการ จาก {MASTER_FILE}")

    def save_master_file(self):
        """บันทึก Master File"""
        try:
            save_data(MASTER_FILE, self.data)
            self.status_var.set(f"บันทึกแล้ว {len(self.data)} รายการ -> {MASTER_FILE}")
            messagebox.showinfo("บันทึกสำเร็จ", f"บันทึกข้อมูล {len(self.data)} รายการแล้ว")
        except Exception as e:
            messagebox.showerror("บันทึกไม่ได้", f"{e}")

    def export_csv(self):
        """Export เป็น CSV ไปที่อื่น"""
        filepath = filedialog.asksaveasfilename(
            title="Export CSV",
            defaultextension=".csv",
            filetypes=[("CSV files", "*.csv"), ("All files", "*.*")],
            initialname=f"ProductMaster_{datetime.now().strftime('%Y%m%d')}.csv",
        )
        if not filepath:
            return
        try:
            save_data(filepath, self.data)
            self.status_var.set(f"Export สำเร็จ -> {filepath}")
            messagebox.showinfo("Export สำเร็จ", f"Export {len(self.data)} รายการ ไปที่:\n{filepath}")
        except Exception as e:
            messagebox.showerror("Export ไม่ได้", f"{e}")

    def export_excel(self):
        """Export เป็น Excel สวยงาม (.xlsx)"""
        if not HAS_OPENPYXL:
            messagebox.showwarning(
                "ต้องติดตั้ง openpyxl",
                "กรุณาเปิด Command Prompt แล้วพิมพ์:\n\npip install openpyxl\n\nแล้วเปิดโปรแกรมใหม่",
            )
            return

        filepath = filedialog.asksaveasfilename(
            title="Export Excel",
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")],
            initialfile=f"ProductMaster_{datetime.now().strftime('%Y%m%d')}.xlsx",
        )
        if not filepath:
            return

        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Product Master"

            # ===== สไตล์ =====
            title_font = Font(name="Tahoma", size=14, bold=True, color="FFFFFF")
            title_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")

            header_font = Font(name="Tahoma", size=10, bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")

            data_font = Font(name="Tahoma", size=10)
            money_font = Font(name="Tahoma", size=10)

            summary_font = Font(name="Tahoma", size=11, bold=True)
            summary_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")

            profit_good_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
            profit_bad_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

            thin_border = Border(
                left=Side(style="thin"),
                right=Side(style="thin"),
                top=Side(style="thin"),
                bottom=Side(style="thin"),
            )
            center_align = Alignment(horizontal="center", vertical="center")
            right_align = Alignment(horizontal="right", vertical="center")
            left_align = Alignment(horizontal="left", vertical="center")

            # ===== Title Row =====
            ws.merge_cells("A1:L1")
            title_cell = ws["A1"]
            title_cell.value = f"Thailand Trophy - Product Master File"
            title_cell.font = title_font
            title_cell.fill = title_fill
            title_cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[1].height = 35

            # ===== Date Row =====
            ws.merge_cells("A2:L2")
            date_cell = ws["A2"]
            date_cell.value = f"Export Date: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
            date_cell.font = Font(name="Tahoma", size=9, italic=True)
            date_cell.alignment = Alignment(horizontal="right")

            # ===== Header Row (Row 3) =====
            headers = [
                "รหัสสินค้า", "ชื่อสินค้า", "หมวดหมู่", "ต้นทุน (บาท)",
                "ราคาขาย (บาท)", "กำไร (บาท)", "Margin %",
                "Supplier", "คงเหลือ", "หน่วย", "หมายเหตุ", "วันที่อัปเดต",
            ]
            col_widths = [14, 30, 14, 14, 14, 14, 12, 18, 10, 8, 20, 14]

            for col_idx, (header, width) in enumerate(zip(headers, col_widths), 1):
                cell = ws.cell(row=3, column=col_idx, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_align
                cell.border = thin_border
                ws.column_dimensions[cell.column_letter].width = width

            ws.row_dimensions[3].height = 25

            # ===== Data Rows =====
            total_cost = 0
            total_price = 0
            data_start_row = 4

            for row_idx, row in enumerate(self.data, data_start_row):
                cost = self._to_float(row.get("ต้นทุน (บาท)", 0))
                price = self._to_float(row.get("ราคาขาย (บาท)", 0))
                profit = price - cost
                margin = (profit / price * 100) if price > 0 else 0

                total_cost += cost
                total_price += price

                values = [
                    row.get("รหัสสินค้า", ""),
                    row.get("ชื่อสินค้า", ""),
                    row.get("หมวดหมู่", ""),
                    cost,
                    price,
                    profit,
                    margin / 100,
                    row.get("Supplier", ""),
                    row.get("จำนวนคงเหลือ", ""),
                    row.get("หน่วย", ""),
                    row.get("หมายเหตุ", ""),
                    row.get("วันที่อัปเดต", ""),
                ]

                for col_idx, val in enumerate(values, 1):
                    cell = ws.cell(row=row_idx, column=col_idx, value=val)
                    cell.font = data_font
                    cell.border = thin_border

                    # จัดตำแหน่ง
                    if col_idx in (1, 3, 9, 10, 12):
                        cell.alignment = center_align
                    elif col_idx in (4, 5, 6):
                        cell.alignment = right_align
                        cell.number_format = '#,##0'
                    elif col_idx == 7:
                        cell.alignment = center_align
                        cell.number_format = '0.0%'
                    else:
                        cell.alignment = left_align

                # สีแถวสลับ
                if row_idx % 2 == 0:
                    stripe_fill = PatternFill(start_color="F2F7FB", end_color="F2F7FB", fill_type="solid")
                    for col_idx in range(1, 13):
                        ws.cell(row=row_idx, column=col_idx).fill = stripe_fill

                # สีกำไร (เขียว/แดง)
                profit_cell = ws.cell(row=row_idx, column=6)
                if profit > 0:
                    profit_cell.fill = profit_good_fill
                elif profit < 0:
                    profit_cell.fill = profit_bad_fill

            # ===== Summary Row =====
            summary_row = data_start_row + len(self.data)
            total_profit = total_price - total_cost
            avg_margin = (total_profit / total_price) if total_price > 0 else 0

            ws.cell(row=summary_row, column=1).value = ""
            ws.cell(row=summary_row, column=2).value = f"TOTAL ({len(self.data)} items)"
            ws.cell(row=summary_row, column=4).value = total_cost
            ws.cell(row=summary_row, column=5).value = total_price
            ws.cell(row=summary_row, column=6).value = total_profit
            ws.cell(row=summary_row, column=7).value = avg_margin

            for col_idx in range(1, 13):
                cell = ws.cell(row=summary_row, column=col_idx)
                cell.font = summary_font
                cell.fill = summary_fill
                cell.border = thin_border
                if col_idx in (4, 5, 6):
                    cell.number_format = '#,##0'
                    cell.alignment = right_align
                elif col_idx == 7:
                    cell.number_format = '0.0%'
                    cell.alignment = center_align

            ws.row_dimensions[summary_row].height = 28

            # ===== Freeze & Filter =====
            ws.auto_filter.ref = f"A3:L{summary_row - 1}"
            ws.freeze_panes = "A4"

            # ===== Save =====
            wb.save(filepath)

            self.status_var.set(f"Export Excel -> {filepath}")
            messagebox.showinfo(
                "Export Excel",
                f"Export {len(self.data)} items\n\n{filepath}",
            )

        except Exception as e:
            messagebox.showerror("Export Error", f"{e}")

    def import_csv(self):
        """Import จาก CSV"""
        filepath = filedialog.askopenfilename(
            title="Import CSV",
            filetypes=[("CSV files", "*.csv"), ("All files", "*.*")],
        )
        if not filepath:
            return

        new_data = load_data(filepath)
        if not new_data:
            messagebox.showwarning("ไม่มีข้อมูล", "ไฟล์ CSV ไม่มีข้อมูลหรืออ่านไม่ได้")
            return

        result = messagebox.askyesnocancel(
            "Import CSV",
            f"พบข้อมูล {len(new_data)} รายการ\n\n"
            f"Yes = แทนที่ข้อมูลเดิมทั้งหมด\n"
            f"No = เพิ่มต่อท้ายข้อมูลเดิม\n"
            f"Cancel = ยกเลิก",
        )

        if result is None:
            return
        elif result:
            self.data = new_data
        else:
            self.data.extend(new_data)

        self.refresh_table()
        self.status_var.set(f"Import สำเร็จ - ตอนนี้มี {len(self.data)} รายการ")

    # ===== Table =====

    def refresh_table(self):
        """แสดงข้อมูลในตาราง (พร้อมกรอง)"""
        self.tree.delete(*self.tree.get_children())

        search = self.search_var.get().lower()
        cat_filter = self.cat_filter_var.get()

        total_cost = 0
        total_price = 0
        shown = 0

        for row in self.data:
            # กรองค้นหา
            if search:
                searchable = f"{row.get('รหัสสินค้า', '')} {row.get('ชื่อสินค้า', '')} {row.get('Supplier', '')}".lower()
                if search not in searchable:
                    continue
            # กรองหมวด
            if cat_filter != "ทั้งหมด" and row.get("หมวดหมู่", "") != cat_filter:
                continue

            cost = self._to_float(row.get("ต้นทุน (บาท)", 0))
            price = self._to_float(row.get("ราคาขาย (บาท)", 0))
            profit = price - cost
            margin = (profit / price * 100) if price > 0 else 0

            self.tree.insert("", "end", values=(
                row.get("รหัสสินค้า", ""),
                row.get("ชื่อสินค้า", ""),
                row.get("หมวดหมู่", ""),
                format_money(cost),
                format_money(price),
                format_money(profit),
                f"{margin:.1f}%",
                row.get("Supplier", ""),
                row.get("จำนวนคงเหลือ", ""),
                row.get("หน่วย", ""),
                row.get("หมายเหตุ", ""),
                row.get("วันที่อัปเดต", ""),
            ))

            total_cost += cost
            total_price += price
            shown += 1

        total_profit = total_price - total_cost
        avg_margin = (total_profit / total_price * 100) if total_price > 0 else 0

        self.summary_var.set(
            f"แสดง {shown}/{len(self.data)} รายการ  |  "
            f"ต้นทุนรวม: {format_money(total_cost)} บาท  |  "
            f"ราคาขายรวม: {format_money(total_price)} บาท  |  "
            f"กำไรรวม: {format_money(total_profit)} บาท  |  "
            f"Margin เฉลี่ย: {avg_margin:.1f}%"
        )

    def sort_column(self, col):
        """เรียงลำดับตารางตามคอลัมน์"""
        reverse = self.sort_reverse.get(col, False)
        items = [(self.tree.set(k, col), k) for k in self.tree.get_children()]

        # พยายามเรียงเป็นตัวเลข
        try:
            items.sort(key=lambda x: float(x[0].replace(",", "").replace("%", "")), reverse=reverse)
        except ValueError:
            items.sort(reverse=reverse)

        for index, (val, k) in enumerate(items):
            self.tree.move(k, "", index)
        self.sort_reverse[col] = not reverse

    # ===== CRUD =====

    def add_product(self):
        """เพิ่มสินค้าใหม่"""
        dialog = ProductDialog(self.root, "เพิ่มสินค้าใหม่")
        if dialog.result:
            self.data.append(dialog.result)
            self.refresh_table()
            self.auto_save()
            self.status_var.set(f"เพิ่มสินค้า: {dialog.result.get('ชื่อสินค้า', '')}")

    def edit_product(self):
        """แก้ไขสินค้า"""
        selection = self.tree.selection()
        if not selection:
            messagebox.showinfo("แจ้งเตือน", "กรุณาเลือกสินค้าที่ต้องการแก้ไข")
            return

        item = self.tree.item(selection[0])
        code = item["values"][0]
        name = item["values"][1]

        # หา row ในข้อมูล
        row_index = None
        for i, row in enumerate(self.data):
            if row.get("รหัสสินค้า", "") == str(code) and row.get("ชื่อสินค้า", "") == str(name):
                row_index = i
                break

        if row_index is None:
            messagebox.showwarning("ไม่พบข้อมูล", "ไม่พบข้อมูลสินค้านี้")
            return

        dialog = ProductDialog(self.root, "แก้ไขสินค้า", self.data[row_index])
        if dialog.result:
            self.data[row_index] = dialog.result
            self.refresh_table()
            self.auto_save()
            self.status_var.set(f"แก้ไขสินค้า: {dialog.result.get('ชื่อสินค้า', '')}")

    def delete_product(self):
        """ลบสินค้า"""
        selection = self.tree.selection()
        if not selection:
            messagebox.showinfo("แจ้งเตือน", "กรุณาเลือกสินค้าที่ต้องการลบ")
            return

        item = self.tree.item(selection[0])
        code = item["values"][0]
        name = item["values"][1]

        result = messagebox.askyesno(
            "ยืนยันการลบ",
            f"ต้องการลบสินค้านี้หรือไม่?\n\n"
            f"รหัส: {code}\n"
            f"ชื่อ: {name}",
        )
        if not result:
            return

        self.data = [
            row for row in self.data
            if not (row.get("รหัสสินค้า", "") == str(code) and row.get("ชื่อสินค้า", "") == str(name))
        ]
        self.refresh_table()
        self.auto_save()
        self.status_var.set(f"ลบสินค้า: {name}")

    def auto_save(self):
        """บันทึกอัตโนมัติทุกครั้งที่มีการเปลี่ยนแปลง"""
        try:
            save_data(MASTER_FILE, self.data)
        except Exception:
            pass

    # ===== Helpers =====

    @staticmethod
    def _to_float(value):
        try:
            return float(str(value).replace(",", ""))
        except (ValueError, TypeError):
            return 0.0


class ProductDialog:
    """Dialog สำหรับเพิ่ม/แก้ไขสินค้า"""

    def __init__(self, parent, title, existing_data=None):
        self.result = None

        self.dialog = tk.Toplevel(parent)
        self.dialog.title(title)
        self.dialog.geometry("500x520")
        self.dialog.resizable(False, False)
        self.dialog.transient(parent)
        self.dialog.grab_set()

        main_frame = ttk.Frame(self.dialog, padding=20)
        main_frame.pack(fill="both", expand=True)

        # ===== ฟิลด์ข้อมูล =====
        fields = [
            ("รหัสสินค้า:", "code"),
            ("ชื่อสินค้า:", "name"),
            ("ต้นทุน (บาท):", "cost"),
            ("ราคาขาย (บาท):", "price"),
            ("Supplier:", "supplier"),
            ("จำนวนคงเหลือ:", "stock"),
            ("หน่วย:", "unit"),
        ]

        self.entries = {}
        row = 0

        for label_text, key in fields:
            ttk.Label(main_frame, text=label_text).grid(row=row, column=0, sticky="w", pady=4)
            entry = ttk.Entry(main_frame, width=35)
            entry.grid(row=row, column=1, sticky="ew", pady=4, padx=(10, 0))
            self.entries[key] = entry
            row += 1

        # หมวดหมู่ (dropdown)
        ttk.Label(main_frame, text="หมวดหมู่:").grid(row=row, column=0, sticky="w", pady=4)
        self.cat_var = tk.StringVar(value=CATEGORIES[0])
        cat_combo = ttk.Combobox(
            main_frame, textvariable=self.cat_var, values=CATEGORIES, state="readonly", width=33
        )
        cat_combo.grid(row=row, column=1, sticky="ew", pady=4, padx=(10, 0))
        row += 1

        # หมายเหตุ (หลายบรรทัด)
        ttk.Label(main_frame, text="หมายเหตุ:").grid(row=row, column=0, sticky="nw", pady=4)
        self.note_text = tk.Text(main_frame, width=35, height=3)
        self.note_text.grid(row=row, column=1, sticky="ew", pady=4, padx=(10, 0))
        row += 1

        main_frame.columnconfigure(1, weight=1)

        # ===== ปุ่ม =====
        btn_frame = ttk.Frame(main_frame)
        btn_frame.grid(row=row, column=0, columnspan=2, pady=15)

        ttk.Button(btn_frame, text="บันทึก", command=self.on_save).pack(side="left", padx=10)
        ttk.Button(btn_frame, text="ยกเลิก", command=self.dialog.destroy).pack(side="left", padx=10)

        # ===== ถ้าแก้ไข ให้ใส่ข้อมูลเดิม =====
        if existing_data:
            self.entries["code"].insert(0, existing_data.get("รหัสสินค้า", ""))
            self.entries["name"].insert(0, existing_data.get("ชื่อสินค้า", ""))
            self.entries["cost"].insert(0, existing_data.get("ต้นทุน (บาท)", ""))
            self.entries["price"].insert(0, existing_data.get("ราคาขาย (บาท)", ""))
            self.entries["supplier"].insert(0, existing_data.get("Supplier", ""))
            self.entries["stock"].insert(0, existing_data.get("จำนวนคงเหลือ", ""))
            self.entries["unit"].insert(0, existing_data.get("หน่วย", ""))
            cat = existing_data.get("หมวดหมู่", "")
            if cat in CATEGORIES:
                self.cat_var.set(cat)
            self.note_text.insert("1.0", existing_data.get("หมายเหตุ", ""))

        self.dialog.wait_window()

    def on_save(self):
        """กดบันทึก"""
        code = self.entries["code"].get().strip()
        name = self.entries["name"].get().strip()

        if not code:
            messagebox.showwarning("กรุณากรอกข้อมูล", "กรุณาใส่รหัสสินค้า", parent=self.dialog)
            return
        if not name:
            messagebox.showwarning("กรุณากรอกข้อมูล", "กรุณาใส่ชื่อสินค้า", parent=self.dialog)
            return

        self.result = {
            "รหัสสินค้า": code,
            "ชื่อสินค้า": name,
            "หมวดหมู่": self.cat_var.get(),
            "ต้นทุน (บาท)": self.entries["cost"].get().strip() or "0",
            "ราคาขาย (บาท)": self.entries["price"].get().strip() or "0",
            "Supplier": self.entries["supplier"].get().strip(),
            "จำนวนคงเหลือ": self.entries["stock"].get().strip() or "0",
            "หน่วย": self.entries["unit"].get().strip() or "ชิ้น",
            "หมายเหตุ": self.note_text.get("1.0", "end-1c").strip(),
            "วันที่อัปเดต": datetime.now().strftime("%Y-%m-%d"),
        }

        self.dialog.destroy()


def main():
    root = tk.Tk()
    ProductMaster(root)
    root.mainloop()


if __name__ == "__main__":
    main()
